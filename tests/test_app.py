import json
from copy import copy
import tempfile
import threading
import unittest
from io import BytesIO
from pathlib import Path
from urllib.request import Request, urlopen
from urllib.error import HTTPError
from http.server import ThreadingHTTPServer
from openpyxl import load_workbook, Workbook
import app
from engine import workbook, columns, read_catalog, suggest, check, export, ean_valid, inspect

class ExcelTests(unittest.TestCase):
    def setUp(self):
        files = app.demo_files(); self.catalog = files['catalogue-demo.xlsx']; self.template = files['template-demo.xlsx']
        self.source, self.rows = read_catalog(self.catalog,'Produits',1)
        self.target = columns(workbook(self.template)['Offres'],2)
        self.mapping = suggest(self.source,self.target)
        self.required = [c['id'] for c in self.target if c['required']]

    def result(self, selected=(2,3)):
        return check(self.rows,self.target,self.mapping,self.required,list(selected))

    def test_suggestions_and_errors(self):
        self.assertEqual(self.mapping, {'1':'1','2':'2','3':'3','4':'4','5':'5','6':'6','7':'7'})
        self.assertEqual(len(self.required),6)
        result = self.result((2,3,4))
        self.assertEqual(len(result['errors']),3)
        self.assertEqual({e['column'] for e in result['errors']},{'EAN *','Description *','Prix *'})
        with self.assertRaises(ValueError): export(self.template,'Offres',2,3,self.target,result)

    def test_export_preserves_workbook(self):
        result = self.result(); self.assertEqual(result['errors'],[])
        data = export(self.template,'Offres',2,3,self.target,result)
        before = workbook(self.template); after = workbook(data)
        self.assertEqual(after.sheetnames,before.sheetnames)
        self.assertEqual(after['Instructions']['A1'].value,before['Instructions']['A1'].value)
        ws=after['Offres']
        self.assertEqual(ws['A3'].value,'DEMO-001');self.assertEqual(ws['B3'].value,'4006381333931')
        self.assertEqual(ws['F4'].value,12.5)
        self.assertEqual(ws['F4'].number_format,before['Offres']['F3'].number_format)
        self.assertEqual(ws['A3'].number_format,before['Offres']['A3'].number_format)
        self.assertEqual(copy(ws['F4'].fill),copy(before['Offres']['F3'].fill))
        self.assertEqual(ws['A2'].comment.text,before['Offres']['A2'].comment.text)
        self.assertEqual(ws.column_dimensions['A'].width,before['Offres'].column_dimensions['A'].width)
        self.assertEqual(ws.freeze_panes,'A3')
        self.assertEqual(len(ws.data_validations.dataValidation),1)
        self.assertIn('F4',ws.data_validations.dataValidation[0])
        self.assertEqual(ws['A5'].value,None)
        # The original upload remains unchanged; a second export has identical cell values.
        self.assertIsNone(workbook(self.template)['Offres']['A3'].value)
        self.assertEqual(workbook(export(self.template,'Offres',2,3,self.target,result))['Offres']['A4'].value,'DEMO-002')

    def test_missing_mapping_and_empty_selection(self):
        self.mapping['4']=''
        self.assertEqual(len(self.result()['errors']),2)
        with self.assertRaises(ValueError): self.result(())

    def test_duplicate_identifiers(self):
        self.rows[1]['values']['1']=self.rows[0]['values']['1']
        self.assertTrue(any('dupliqué' in e['message'] for e in self.result()['errors']))

    def test_do_not_overwrite_existing_or_merged_cells(self):
        wb=workbook(self.template); wb['Offres']['A3']='instruction'; data=BytesIO();wb.save(data)
        with self.assertRaisesRegex(ValueError,'contient déjà'): export(data.getvalue(),'Offres',2,3,self.target,self.result())
        wb=workbook(self.template);wb['Offres'].merge_cells('A3:B3');data=BytesIO();wb.save(data)
        with self.assertRaisesRegex(ValueError,'fusionnées'): export(data.getvalue(),'Offres',2,3,self.target,self.result())

    def test_ean_and_leading_zero(self):
        self.assertTrue(ean_valid('4006381333931'));self.assertFalse(ean_valid('1234567890123'))
        wb=Workbook();ws=wb.active;ws.append(['SKU','EAN']);ws.append([123,'0001234567895']);ws['A2'].number_format='000000'
        stream=BytesIO();wb.save(stream);source,rows=read_catalog(stream.getvalue(),ws.title,1)
        self.assertEqual(rows[0]['values']['1'],'000123');self.assertEqual(rows[0]['values']['2'],'0001234567895')

    def test_formula_and_ambiguous_alias(self):
        wb=Workbook();ws=wb.active;ws.append(['SKU']);ws.append(['=1+1']);stream=BytesIO();wb.save(stream)
        with self.assertRaisesRegex(ValueError,'formule'):read_catalog(stream.getvalue(),ws.title,1)
        self.assertEqual(suggest([{'id':'1','label':'EAN','kind':'ean'},{'id':'2','label':'EAN','kind':'ean'}],[{'id':'1','label':'EAN *','kind':'ean'}]),{'1':''})

class HTTPTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp=tempfile.TemporaryDirectory();cls.old_profiles=app.PROFILES;app.PROFILES=Path(cls.temp.name)/'mappings.json'
        cls.server=ThreadingHTTPServer(('127.0.0.1',0),app.Handler)
        cls.thread=threading.Thread(target=cls.server.serve_forever,daemon=True);cls.thread.start()
        cls.base=f'http://127.0.0.1:{cls.server.server_port}'
    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown();cls.server.server_close();app.PROFILES=cls.old_profiles;cls.temp.cleanup()
    def setUp(self):
        self.token=json.load(urlopen(self.base+'/api/session'))['token']
    def post(self,path,payload):
        return urlopen(Request(self.base+path,data=json.dumps(payload).encode(),headers={'Content-Type':'application/json','X-Session':self.token}))
    def upload(self,role,data,name):
        boundary='test-boundary'
        raw=(f'--{boundary}\r\nContent-Disposition: form-data; name="role"\r\n\r\n{role}\r\n--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="{name}"\r\nContent-Type: application/octet-stream\r\n\r\n').encode()+data+f'\r\n--{boundary}--\r\n'.encode()
        return json.load(urlopen(Request(self.base+'/api/upload',data=raw,headers={'X-Session':self.token,'Content-Type':f'multipart/form-data; boundary={boundary}'})))
    def test_full_import_control_export_and_profiles(self):
        files=app.demo_files()
        for role,name in [('catalog','catalogue-demo.xlsx'),('template','template-demo.xlsx')]:
            response=self.upload(role,files[name],name);self.assertTrue(response['sheets'])
        prepared=json.load(self.post('/api/prepare',{'catalogSheet':'Produits','catalogHeader':1,'templateSheet':'Offres','templateHeader':2,'startRow':3}))
        payload={'mapping':prepared['mapping'],'required':[c['id'] for c in prepared['target'] if c['required']],'selected':[2,3,4]}
        self.assertEqual(len(json.load(self.post('/api/check',payload))['errors']),3)
        with self.assertRaises(HTTPError) as ctx:self.post('/api/export',payload)
        self.assertEqual(ctx.exception.code,400)
        payload['selected']=[2,3]
        self.assertEqual(json.load(self.post('/api/check',payload))['errors'],[])
        output=self.post('/api/export',payload).read()
        self.assertEqual(workbook(output)['Offres']['A4'].value,'DEMO-002')
        profile={'mapping':prepared['mapping'],'required':payload['required'],'signature':'test'}
        saved=json.load(self.post('/api/profile',{'name':'Test marketplace','profile':profile}))
        self.assertEqual(saved['profiles']['Test marketplace'],profile)
        next_session=json.load(urlopen(self.base+'/api/session'))
        self.assertEqual(next_session['profiles']['Test marketplace'],profile)
        # New session does not inherit product uploads.
        self.token=next_session['token']
        with self.assertRaises(HTTPError):self.post('/api/check',payload)

    def test_reject_bad_files_and_origin(self):
        with self.assertRaises(HTTPError):self.upload('catalog',b'invalid','bad.xlsx')
        with self.assertRaises(HTTPError):self.upload('catalog',b'invalid','bad.xlsm')
        with self.assertRaises(HTTPError) as ctx:
            urlopen(Request(self.base+'/api/check',data=b'{}',headers={'X-Session':self.token,'Origin':'http://other-site.example'}))
        self.assertEqual(ctx.exception.code,403)

if __name__=='__main__':unittest.main()
