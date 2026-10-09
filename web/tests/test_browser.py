"""Functional browser tests with synthetic data only. No application backend."""
import base64
from copy import copy
import json
import os
import tempfile
import threading
import unittest
from pathlib import Path
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED
from functools import partial
from http.server import ThreadingHTTPServer, SimpleHTTPRequestHandler
from xml.etree import ElementTree as ET
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

WEB = Path(__file__).resolve().parents[1]
DEMOS = json.loads((WEB / 'demo-data.js').read_text().removeprefix('window.SEMIN_DEMO=').strip().removesuffix(';'))
S = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'

class QuietHandler(SimpleHTTPRequestHandler):
    def log_message(self, *args): pass

class BrowserTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.server = ThreadingHTTPServer(('127.0.0.1', 0), partial(QuietHandler, directory=str(WEB)))
        cls.thread = threading.Thread(target=cls.server.serve_forever, daemon=True); cls.thread.start()
        cls.url = os.environ.get('SEMIN_WEB_TEST_URL') or f'http://127.0.0.1:{cls.server.server_port}/SEMIN-Marketplace.html'
        cls.playwright = sync_playwright().start()
        args = {'headless': True}
        if os.environ.get('SEMIN_TEST_CHROMIUM'):
            args['executable_path'] = os.environ['SEMIN_TEST_CHROMIUM']
        cls.browser = cls.playwright.chromium.launch(**args)
    @classmethod
    def tearDownClass(cls):
        cls.browser.close(); cls.playwright.stop(); cls.server.shutdown(); cls.server.server_close(); cls.temp.cleanup()
    def setUp(self):
        self.context = self.browser.new_context(viewport={'width': 820, 'height': 1180}, is_mobile=True, has_touch=True, locale='fr-FR')
        self.page = self.context.new_page()
        self.errors = []; self.requests = []
        self.page.on('pageerror', lambda error: self.errors.append(str(error)))
        self.page.on('request', lambda request: self.requests.append(request.url))
        self.page.goto(self.url)
    def tearDown(self):
        self.assertFalse(self.errors)
        self.context.close()
    def demo(self):
        self.page.get_by_role('button', name='Charger la démonstration').tap()
        self.page.wait_for_selector('#stat-errors:text-is("3")')
    def valid(self):
        self.page.get_by_role('checkbox', name='Sélectionner la ligne 4', exact=True).uncheck()
        self.page.get_by_role('button', name='Contrôler les données').tap()
        self.page.wait_for_selector('#stat-errors:text-is("0")')
    def download(self):
        with self.page.expect_download() as event:
            self.page.get_by_role('button', name='Exporter le fichier Excel').tap()
        path = Path(self.temp.name) / 'browser-export.xlsx'; event.value.save_as(path)
        return path
    def test_touch_demo_validation_export_and_no_network(self):
        self.demo(); self.assertTrue(self.page.locator('#export').is_disabled())
        self.valid(); path = self.download(); book = load_workbook(path)
        self.assertEqual(book.sheetnames, ['Offres', 'Instructions'])
        ws = book['Offres']
        self.assertEqual(ws['A3'].value, 'DEMO-001'); self.assertEqual(ws['B4'].value, '5901234123457')
        self.assertEqual(ws['F4'].value, 12.5); self.assertEqual(ws['F4'].number_format, '#,##0.00')
        self.assertEqual(copy(ws['A3'].fill), copy(ws['A4'].fill))
        self.assertEqual(ws.freeze_panes, 'A3'); self.assertTrue(ws['A2'].comment)
        self.assertIn('F4', ws.data_validations.dataValidation[0])
        with ZipFile(BytesIO(base64.b64decode(DEMOS['template-demo.xlsx']))) as original, ZipFile(path) as exported:
            self.assertEqual(set(original.namelist()), set(exported.namelist()))
            self.assertEqual([name for name in original.namelist() if original.read(name) != exported.read(name)], ['xl/worksheets/sheet1.xml'])
        self.assertEqual(self.requests, [self.url])  # No API calls, uploads, CDNs or remote resources.
        self.assertIn('2 sélectionné(s)', self.page.locator('#product-count').inner_text())
        self.page.set_viewport_size({'width': 390, 'height': 844})
        self.assertTrue(self.page.evaluate('document.documentElement.scrollWidth <= innerWidth'))
    def test_mapping_persistence_backup_restore_and_filtered_selection(self):
        self.demo()
        self.page.get_by_label('Source pour Titre *', exact=True).select_option('')
        self.assertTrue(self.page.locator('#export').is_disabled())
        self.page.locator('#marketplace').fill('Marketplace fictive test')
        self.page.get_by_role('button', name='Enregistrer les correspondances', exact=True).tap()
        self.page.wait_for_selector('#notice:text-is("Correspondances enregistrées localement pour cette marketplace.")')
        self.page.reload(); self.demo()
        self.page.locator('#profiles').select_option('Marketplace fictive test')
        self.assertEqual(self.page.get_by_label('Source pour Titre *', exact=True).input_value(), '')
        self.page.get_by_role('button', name='Contrôler les données').tap()
        self.page.wait_for_selector('#stat-errors:text-is("6")')
        with self.page.expect_download() as event:
            self.page.get_by_role('button', name='Sauvegarder les profils JSON').tap()
        backup = Path(self.temp.name) / 'profiles.json'; event.value.save_as(backup)
        profiles = json.loads(backup.read_text()); self.assertIn('Marketplace fictive test', profiles)
        self.assertNotIn('DEMO-001', backup.read_text())
        self.page.evaluate('localStorage.clear()'); self.page.reload(); self.demo()
        self.page.locator('#restore-profiles').set_input_files(backup)
        self.page.wait_for_selector('#notice:text-is("Correspondances importées sur cet appareil. Les profils de même nom sont remplacés.")')
        self.page.locator('#profiles').select_option('Marketplace fictive test')
        self.assertEqual(self.page.get_by_label('Source pour Titre *', exact=True).input_value(), '')
        self.page.get_by_label('Source pour Titre *', exact=True).select_option('3')
        self.page.get_by_role('button', name='Tout désélectionner').tap()
        self.page.locator('#search').fill('4006381333931')
        self.page.get_by_role('button', name='Sélectionner les résultats').tap()
        self.page.get_by_role('button', name='Contrôler les données').tap()
        self.page.wait_for_selector('#stat-errors:text-is("0")')
        book = load_workbook(self.download()); self.assertEqual(book['Offres']['A3'].value, 'DEMO-001'); self.assertIsNone(book['Offres']['A4'].value)
    def test_real_file_import_and_shared_strings(self):
        # Exercise the alternate shared-string representation used by many Excel catalogues.
        raw = base64.b64decode(DEMOS['catalogue-demo.xlsx'])
        with ZipFile(BytesIO(raw)) as archive:
            parts = {name: archive.read(name) for name in archive.namelist()}
        doc = ET.fromstring(parts['xl/worksheets/sheet1.xml'])
        for cell in doc.iter(f'{{{S}}}c'):
            if cell.get('r') == 'A2':
                for child in list(cell): cell.remove(child)
                cell.set('t', 's'); ET.SubElement(cell, f'{{{S}}}v').text = '0'
        parts['xl/worksheets/sheet1.xml'] = ET.tostring(doc)
        parts['xl/sharedStrings.xml'] = f'<sst xmlns="{S}" count="1" uniqueCount="1"><si><t>DEMO-001</t></si></sst>'.encode()
        rel = ET.fromstring(parts['xl/_rels/workbook.xml.rels']); rel_ns = 'http://schemas.openxmlformats.org/package/2006/relationships'
        ET.SubElement(rel, f'{{{rel_ns}}}Relationship', Id='rIdSharedTest', Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings', Target='sharedStrings.xml')
        parts['xl/_rels/workbook.xml.rels'] = ET.tostring(rel)
        ct = ET.fromstring(parts['[Content_Types].xml'])
        ET.SubElement(ct, '{http://schemas.openxmlformats.org/package/2006/content-types}Override', PartName='/xl/sharedStrings.xml', ContentType='application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml')
        parts['[Content_Types].xml'] = ET.tostring(ct)
        catalog = Path(self.temp.name) / 'shared-catalog.xlsx'
        with ZipFile(catalog, 'w', ZIP_DEFLATED) as archive:
            for name, data in parts.items(): archive.writestr(name, data)
        template = Path(self.temp.name) / 'template.xlsx'; template.write_bytes(base64.b64decode(DEMOS['template-demo.xlsx']))
        for role, file in [('catalog', catalog), ('template', template)]:
            self.page.locator(f'#{role}-file').set_input_files(file)
            self.page.wait_for_selector(f'#{role}-name:text-is("{file.name}")')
        self.page.locator('#template-header').fill('2'); self.page.locator('#start-row').fill('3')
        self.page.get_by_role('button', name='Analyser les fichiers').tap()
        self.page.wait_for_selector('#stat-products:text-is("3")')
        self.page.locator('#search').fill('DEMO-001'); self.page.get_by_role('button', name='Sélectionner les résultats').tap()
        self.page.get_by_role('button', name='Contrôler les données').tap(); self.page.wait_for_selector('#stat-errors:text-is("0")')
        self.assertEqual(load_workbook(self.download())['Offres']['A3'].value, 'DEMO-001')
    def test_export_refuses_existing_or_merged_destination(self):
        for merged in (False, True):
            self.page.reload(); self.demo()
            book = load_workbook(BytesIO(base64.b64decode(DEMOS['template-demo.xlsx'])))
            if merged: book['Offres'].merge_cells('A3:B3')
            else: book['Offres']['A3'] = 'Instructions à conserver'
            template = Path(self.temp.name) / 'occupied.xlsx'; book.save(template)
            self.page.locator('#template-file').set_input_files(template)
            self.page.wait_for_selector('#template-name:text-is("occupied.xlsx")')
            self.page.get_by_role('button', name='Analyser les fichiers').tap()
            self.page.wait_for_selector('#stat-products:text-is("3")')
            self.page.locator('#search').fill('DEMO-001'); self.page.get_by_role('button', name='Sélectionner les résultats').tap()
            self.page.get_by_role('button', name='Contrôler les données').tap(); self.page.wait_for_selector('#stat-errors:text-is("0")')
            self.page.get_by_role('button', name='Exporter le fichier Excel').tap()
            self.page.wait_for_selector('#notice.error')
            message = self.page.locator('#notice').inner_text()
            self.assertIn('fusionnées' if merged else 'déjà une valeur', message)
            self.assertTrue(self.page.locator('#export').is_disabled())
    def test_no_product_data_in_browser_storage(self):
        self.demo(); self.valid()
        self.page.locator('#marketplace').fill('Exemple')
        self.page.get_by_role('button', name='Enregistrer les correspondances', exact=True).tap()
        self.page.wait_for_selector('#notice:text-is("Correspondances enregistrées localement pour cette marketplace.")')
        saved = self.page.evaluate('JSON.stringify(localStorage)')
        self.assertNotIn('DEMO-001', saved); self.assertNotIn('4006381333931', saved)
        self.assertIn('Exemple', saved)

if __name__ == '__main__': unittest.main()
