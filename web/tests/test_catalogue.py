import json
from pathlib import Path
from io import BytesIO
from openpyxl import Workbook
import unittest
import test_browser as base

class CatalogueTests(unittest.TestCase):
    setUpClass=classmethod(base.BrowserTests.setUpClass.__func__)
    tearDownClass=classmethod(base.BrowserTests.tearDownClass.__func__)
    setUp=base.BrowserTests.setUp
    tearDown=base.BrowserTests.tearDown

    def catalogue(self):
        self.page.locator('[data-studio-view="catalogue"]').tap()
    def demo_import(self):
        self.catalogue()
        self.page.locator('#master-demo').tap()
        self.page.wait_for_selector('#master-import-plan:not([hidden])')
        self.page.locator('#master-apply').tap()
        self.page.wait_for_selector('#master-total:text-is("3")',state='attached')
    def edit(self,sku):
        self.page.get_by_role('button',name='Ouvrir la fiche '+sku,exact=True).tap()
        self.page.wait_for_selector('#master-editor[open]')
    def write_catalog(self,name,rows):
        wb=Workbook();ws=wb.active;ws.title='Catalogue fictif';ws.append(['SKU','EAN','Désignation','Marque','Famille']);
        for row in rows:ws.append(row)
        path=Path(self.temp.name)/name;wb.save(path)
        self.catalogue();self.page.locator('#master-file').set_input_files(path)
        self.page.wait_for_selector(f'#master-file-name:text-is("{name}")')
        self.page.locator('#master-read').tap();self.page.wait_for_selector('#master-plan:not([hidden])')
        self.page.locator('#master-plan').tap();self.page.wait_for_selector('#master-import-plan:not([hidden])')
    def test_master_import_edit_search_and_indexeddb_reload(self):
        self.demo_import()
        self.assertEqual(self.page.locator('#master-complete').text_content(),'2')
        self.assertEqual(self.page.locator('#master-incomplete').text_content(),'1')
        self.page.locator('#master-plan').tap();self.page.wait_for_selector('#master-import-plan:not([hidden])')
        self.assertIn('Identique : 3',self.page.locator('#master-plan-summary').text_content())
        self.page.locator('#master-cancel').tap()
        self.edit('DEMO-001')
        self.page.get_by_label('Famille',exact=True).last.fill('Famille fictive')
        self.page.get_by_label('Désignation',exact=True).fill('Désignation fictive modifiée')
        self.page.get_by_label('URLs images',exact=True).fill('https://example.invalid/image-fictive.png')
        self.page.locator('#master-editor-save').tap();self.page.wait_for_selector('#master-editor',state='hidden')
        self.page.reload();self.page.wait_for_selector('#master-total:text-is("3")')
        self.catalogue();self.page.locator('#master-search').fill('Famille fictive')
        self.assertEqual(self.page.locator('#master-products tbody tr').count(),1)
        self.assertIn('Désignation fictive modifiée',self.page.locator('#master-products').text_content())
        self.page.locator('#master-select').tap();self.assertIn('1 sélectionné(s)',self.page.locator('#master-list-count').text_content())
        self.assertEqual(self.requests,[self.url,self.url])
    def test_reimport_requires_explicit_selection_and_shows_erasure(self):
        self.demo_import()
        self.write_catalog('changes.xlsx',[
            ['DEMO-001','4006381333931','Désignation revue',None,'Famille test'],
            ['DEMO-004','', 'Nouveau produit fictif','Marque fictive','Famille test']])
        update=self.page.get_by_role('checkbox',name='Appliquer la ligne 2',exact=True)
        self.assertFalse(update.is_checked())
        self.assertIn('Marque : SEMIN Démo → ∅',self.page.locator('#master-plan-table').text_content())
        self.page.locator('#master-apply').tap();self.page.wait_for_selector('#master-total:text-is("4")',state='attached')
        self.edit('DEMO-001');self.assertEqual(self.page.get_by_label('Marque',exact=True).input_value(),'SEMIN Démo');self.page.locator('#master-editor-close').tap()
        self.write_catalog('changes-confirmed.xlsx',[['DEMO-001','4006381333931','Désignation revue',None,'Famille test']])
        self.page.get_by_role('checkbox',name='Appliquer la ligne 2',exact=True).check()
        self.page.locator('#master-apply').tap();self.page.wait_for_selector('#master-import-plan',state='hidden')
        self.edit('DEMO-001');self.assertEqual(self.page.get_by_label('Marque',exact=True).input_value(),'')
        self.assertIn('Marque',self.page.locator('#master-editor-issues').text_content())
    def test_duplicates_are_blocked_and_edit_cannot_create_duplicate(self):
        self.write_catalog('duplicates.xlsx',[
            ['FAKE-1','', 'Fictif 1','Fictif',''],['FAKE-1','','Fictif 2','Fictif',''],
            ['FAKE-2','5901234123457','Fictif 3','Fictif',''],['FAKE-3','5901234123457','Fictif 4','Fictif','']])
        self.assertIn('Conflit — bloqué : 4',self.page.locator('#master-plan-summary').text_content())
        self.assertEqual(self.page.locator('#master-plan-table input:disabled').count(),4)
        self.page.locator('#master-apply').tap();self.page.wait_for_selector('#notice.error')
        self.assertEqual(self.page.locator('#master-total').text_content(),'0')
        self.demo_import();self.edit('DEMO-001');self.page.get_by_label('SKU',exact=True).fill('DEMO-002')
        self.page.locator('#master-editor-save').tap();self.page.wait_for_selector('#master-editor-message:text-is("Ce SKU ou cet EAN est déjà utilisé par une autre fiche.")')
        self.page.locator('#master-editor-close').tap();self.assertIn('DEMO-001',self.page.locator('#master-products').text_content())
    def test_custom_attribute_backup_restore_and_invalid_backup_atomicity(self):
        self.demo_import();self.page.locator('[data-studio-view="parameters"]').tap()
        self.page.locator('#attribute-name').fill('Attribut technique fictif');self.page.locator('#attribute-type').select_option('number');self.page.locator('#attribute-add').tap()
        self.page.wait_for_selector('#notice:text-is("Attribut ajouté. Il est disponible dans les fiches et les prochains imports.")')
        self.catalogue();self.edit('DEMO-001');self.page.get_by_label('Attribut technique fictif',exact=True).fill('42')
        self.page.locator('#master-editor-save').tap();self.page.wait_for_selector('#master-editor',state='hidden')
        self.page.locator('[data-studio-view="parameters"]').tap()
        with self.page.expect_download() as event:self.page.locator('#master-backup').tap()
        path=Path(self.temp.name)/'catalogue-backup.json';event.value.save_as(path)
        backup=json.loads(path.read_text());self.assertEqual(len(backup['catalogue']['products']),3)
        self.assertIn('Attribut technique fictif',[a['label'] for a in backup['catalogue']['attributes']])
        # Edits after the backup must not be overwritten until the replacement is explicitly confirmed.
        self.catalogue();self.edit('DEMO-001');self.page.get_by_label('Attribut technique fictif',exact=True).fill('43');self.page.locator('#master-editor-save').tap();self.page.wait_for_selector('#master-editor',state='hidden')
        self.page.locator('[data-studio-view="parameters"]').tap();self.page.locator('#master-restore-file').set_input_files(path)
        self.page.wait_for_selector('#master-restore-preview:not([hidden])');self.assertTrue(self.page.locator('#master-restore').is_disabled())
        self.page.locator('#master-restore-consent').check();self.page.locator('#master-restore').tap();self.page.wait_for_selector('#notice:text-is("Sauvegarde restaurée localement.")')
        self.catalogue();self.edit('DEMO-001');self.assertEqual(self.page.get_by_label('Attribut technique fictif',exact=True).input_value(),'42');self.page.locator('#master-editor-close').tap()
        invalid=Path(self.temp.name)/'invalid-backup.json';invalid.write_text('{"format":"SEMIN-MASTER-CATALOGUE","version":1,"catalogue":{}}')
        self.page.locator('[data-studio-view="parameters"]').tap();self.page.locator('#master-restore-file').set_input_files(invalid);self.page.wait_for_selector('#notice.error')
        self.assertEqual(self.page.locator('#master-total').text_content(),'3')
    def test_cross_tab_revision_prevents_lost_updates(self):
        self.demo_import();other=self.context.new_page();other.goto(self.url);other.wait_for_selector('#master-total:text-is("3")')
        other.locator('[data-studio-view="catalogue"]').tap();other.get_by_role('button',name='Ouvrir la fiche DEMO-001',exact=True).tap()
        self.edit('DEMO-001');self.page.get_by_label('Désignation',exact=True).fill('Modification onglet 1');self.page.locator('#master-editor-save').tap();self.page.wait_for_selector('#master-editor',state='hidden')
        other.get_by_label('Désignation',exact=True).fill('Modification onglet 2');other.locator('#master-editor-save').tap()
        other.wait_for_selector('#master-editor-message:text-is("Le catalogue a été modifié dans un autre onglet. Actualisez les données puis recommencez.")');other.close()
        self.page.reload();self.page.wait_for_selector('#master-total:text-is("3")');self.catalogue();self.edit('DEMO-001');self.assertEqual(self.page.get_by_label('Désignation',exact=True).input_value(),'Modification onglet 1')
    def test_custom_sku_mapping_preserves_numeric_leading_zero_format(self):
        wb=Workbook();ws=wb.active;ws.append(['Code interne','Désignation','Marque']);ws.append([123,'Produit fictif','Marque fictive']);ws['A2'].number_format='000000'
        path=Path(self.temp.name)/'leading-zero-master.xlsx';wb.save(path)
        self.catalogue();self.page.locator('#master-file').set_input_files(path);self.page.wait_for_selector('#master-file-name:text-is("leading-zero-master.xlsx")')
        self.page.locator('#master-read').tap();self.page.wait_for_selector('#master-plan:not([hidden])')
        self.page.get_by_label('Attribut maître pour Code interne',exact=True).select_option('sku')
        self.page.locator('#master-plan').tap();self.page.wait_for_selector('#master-import-plan:not([hidden])')
        self.assertIn('000123',self.page.locator('#master-plan-table').text_content())
        self.page.locator('#master-apply').tap();self.page.wait_for_selector('#master-total:text-is("1")',state='attached')
        self.edit('000123');self.assertEqual(self.page.get_by_label('SKU',exact=True).input_value(),'000123')

    def test_corrupt_existing_storage_is_not_silently_overwritten(self):
        self.demo_import()
        self.page.evaluate("""async()=>{const db=await new Promise((resolve,reject)=>{const req=indexedDB.open('semin-marketplace-studio',1);req.onsuccess=()=>resolve(req.result);req.onerror=()=>reject(req.error);});await new Promise((resolve,reject)=>{const tx=db.transaction('workspace','readwrite');tx.objectStore('workspace').put({schemaVersion:1,revision:0,products:'corrupted-test-marker'},'catalogue');tx.oncomplete=resolve;tx.onerror=reject;});db.close();}""")
        self.page.reload();self.page.wait_for_selector('#notice.error');self.catalogue()
        self.page.locator('#master-demo').tap();self.page.wait_for_selector('#master-import-plan:not([hidden])');self.page.locator('#master-apply').tap()
        self.page.wait_for_selector('#notice:text-is("Le catalogue local n’a pas pu être chargé. Actualisez les données avant tout enregistrement.")')
        stored=self.page.evaluate("""async()=>{const db=await new Promise(resolve=>{const req=indexedDB.open('semin-marketplace-studio',1);req.onsuccess=()=>resolve(req.result);});const value=await new Promise(resolve=>{const req=db.transaction('workspace','readonly').objectStore('workspace').get('catalogue');req.onsuccess=()=>resolve(req.result);});db.close();return value;}""")
        self.assertEqual(stored['products'],'corrupted-test-marker')

    def test_mobile_master_layout_and_no_network_after_import(self):
        self.demo_import();self.page.set_viewport_size({'width':390,'height':844})
        self.assertTrue(self.page.evaluate('document.documentElement.scrollWidth <= innerWidth'))
        self.edit('DEMO-001');self.assertTrue(self.page.evaluate('document.documentElement.scrollWidth <= innerWidth'))
        self.assertEqual(self.requests,[self.url])

if __name__=='__main__':unittest.main()
