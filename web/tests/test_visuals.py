"""URL-only visual references. Synthetic URLs, no image hosting or real product data."""
import base64
from pathlib import Path
import unittest
from openpyxl import load_workbook, Workbook
import test_onebase_castorama as base

class VisualTests(unittest.TestCase):
    setUpClass=classmethod(base.OneBaseCastoramaTests.setUpClass.__func__)
    tearDownClass=classmethod(base.OneBaseCastoramaTests.tearDownClass.__func__)
    setUp=base.OneBaseCastoramaTests.setUp
    tearDown=base.OneBaseCastoramaTests.tearDown
    nav=base.OneBaseCastoramaTests.nav
    analyse=base.OneBaseCastoramaTests.analyse
    source=base.OneBaseCastoramaTests.source
    save=base.OneBaseCastoramaTests.save
    product=base.OneBaseCastoramaTests.product
    template=base.OneBaseCastoramaTests.template
    upload=base.OneBaseCastoramaTests.upload
    check=base.OneBaseCastoramaTests.check
    download=base.OneBaseCastoramaTests.download
    def open_visuals(self):
        self.nav('visuals');self.page.wait_for_function("() => document.querySelector('#visual-product').options.length>1")
        self.page.locator('#visual-product').select_option(index=1)
    def field(self,slot):return self.page.locator('[data-visual-slot="'+slot+'"]')
    def save_visuals(self):
        self.page.locator('#visual-plan').tap();self.page.wait_for_selector('#visual-confirm:not([hidden])');self.page.locator('#visual-consent').check();self.page.locator('#visual-save').tap();self.page.wait_for_function("() => document.querySelector('#notice').textContent==='URL des visuels enregistrées localement.'")
    def test_url_validation_admin_and_batch_duplicates(self):
        result=self.page.evaluate('''() => ({bad:['data:image/png,abc','file:///image.png','javascript:alert(1)','https://user:secret@example.invalid/a','https://example.invalid/a b','https://semin-produits.onebase.fr/ADMIN/openfile.php','https://semin-produits.onebase.fr/%41DMIN/image.jpg'].map(Visuals.issue),good:Visuals.issue('https://example.invalid/image.jpg?v=1'), duplicate:(()=>{try{Visuals.batch('MAIN: https://example.invalid/1\nMAIN: https://example.invalid/2',{});return false}catch{return true}})()})'''.replace('/1\nMAIN','/1\\nMAIN'))
        self.assertTrue(all(result['bad']));self.assertEqual(result['good'],'');self.assertTrue(result['duplicate'])
    def test_batch_order_extra_slots_confirmation_and_json(self):
        self.product();before=self.page.evaluate('CatalogueStore.read()');self.open_visuals()
        self.page.locator('#visual-batch').fill('https://example.invalid/main.jpg\nhttps://example.invalid/second.jpg\nPT07: https://example.invalid/extra.jpg')
        self.page.locator('#visual-batch-apply').tap();self.assertEqual(self.field('PT07').input_value(),'https://example.invalid/extra.jpg')
        self.field('PT01').locator('../..').get_by_role('button',name='Définir comme principale').tap()
        self.assertEqual(self.field('MAIN').input_value(),'https://example.invalid/second.jpg')
        self.assertEqual(self.page.evaluate('CatalogueStore.read()'),before)
        self.page.locator('#visual-plan').tap();self.assertTrue(self.page.locator('#visual-save').is_disabled())
        self.save_visuals();after=self.page.evaluate('CatalogueStore.read()')
        self.assertEqual(after['products'][0]['dossier'],before['products'][0]['dossier'])
        self.assertEqual(after['products'][0]['values']['visual_MAIN'],'https://example.invalid/second.jpg')
        backup=self.page.evaluate('CatalogueStore.read().then(CatalogueModel.backup)')
        restored=self.page.evaluate('(backup)=>CatalogueModel.restore(CatalogueModel.empty(),backup)',backup)
        self.assertEqual(restored['products'][0]['values']['visual_PT07'],'https://example.invalid/extra.jpg')
        self.page.reload();self.open_visuals();self.assertEqual(self.field('MAIN').input_value(),'https://example.invalid/second.jpg')
        self.assertFalse(any(r.startswith('https://example.invalid') for r in self.requests))
    def test_explicit_reuse_preserves_sources(self):
        self.product();self.page.evaluate('''async()=>{const d=await CatalogueStore.read();const p=structuredClone(d.products[0]);p.id=CatalogueModel.uid();p.values.sku='FICTIF-002';p.values.ean='';d.products.push(p);await CatalogueStore.save(d,d.revision)}''')
        self.open_visuals();before=self.page.evaluate('CatalogueStore.read()')
        self.field('MAIN').fill('https://example.invalid/common.jpg');self.page.locator('#visual-reuse').select_option(index=1)
        self.save_visuals();after=self.page.evaluate('CatalogueStore.read()')
        for index,p in enumerate(after['products']):
            self.assertEqual(p['values']['visual_MAIN'],'https://example.invalid/common.jpg');self.assertEqual(p['dossier'],before['products'][index]['dossier'])
    def test_preview_success_failure_and_admin_do_not_send_catalogue(self):
        self.product();self.open_visuals();url='https://assets.example.invalid/pixel.png'
        png=base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jD9kAAAAASUVORK5CYII=')
        self.page.route(url,lambda route:route.fulfill(status=200,content_type='image/png',body=png))
        self.field('MAIN').fill(url);card=self.field('MAIN').locator('../..');card.get_by_role('button',name='Afficher / contrôler').tap()
        self.page.wait_for_function("() => document.querySelector('.visual-status').textContent.includes('Image chargée')")
        self.assertTrue(card.locator('img').is_visible())
        failed='https://assets.example.invalid/missing.png';self.page.route(failed,lambda route:route.abort())
        self.field('MAIN').fill(failed);card.get_by_role('button',name='Afficher / contrôler').tap()
        self.page.wait_for_function("() => document.querySelector('.visual-status').textContent.includes('non concluant')")
        admin='https://semin-produits.onebase.fr/ADMIN/image.jpg';self.field('MAIN').fill(admin);card.get_by_role('button',name='Afficher / contrôler').tap()
        self.assertNotIn(admin,self.requests)
        self.assertEqual([r for r in self.requests if r.startswith('https:')],[url,failed])
    def test_castorama_url_export_and_structure(self):
        self.product();self.open_visuals();self.field('MAIN').fill('https://example.invalid/main.jpg');self.field('PT01').fill('https://example.invalid/detail.jpg');self.save_visuals()
        path=self.template(large=False);wb=load_workbook(path);wb['Data']['J1']='MAIN';wb['Data']['K1']='other_image_url1';wb.save(path)
        self.upload(path)
        self.assertEqual(self.page.get_by_label('Mapping Castorama · MAIN',exact=True).input_value(),'master.visual_MAIN')
        self.assertEqual(self.page.get_by_label('Mapping Castorama · other_image_url1',exact=True).input_value(),'master.visual_PT01')
        self.check();result=load_workbook(self.download())
        self.assertEqual(result['Data']['J3'].value,'https://example.invalid/main.jpg');self.assertEqual(result['Data']['K3'].value,'https://example.invalid/detail.jpg')
        self.assertEqual(result.sheetnames,wb.sheetnames);self.assertEqual(list(result['Data'].values)[1],list(wb['Data'].values)[1])
        self.page.evaluate('''async()=>{const d=await CatalogueStore.read();d.products[0].values.visual_MAIN='https://semin-produits.onebase.fr/ADMIN/image.jpg';await CatalogueStore.save(d,d.revision)}''')
        wb['Data']['J1']='Photo spéciale';special=path.with_name('template-visuels-personnalise.xlsx');wb.save(special);self.upload(special)
        self.page.get_by_label('Mapping Castorama · Photo spéciale',exact=True).select_option('master.visual_MAIN');self.check()
        self.assertIn('administration OneBase',self.page.locator('#cast-errors').text_content());self.assertTrue(self.page.locator('#cast-export').is_disabled())
    def test_stale_edit_is_blocked_and_mobile_layout(self):
        self.product();self.open_visuals();self.field('MAIN').fill('https://example.invalid/image.jpg')
        self.page.evaluate('CatalogueStore.read().then(d=>CatalogueStore.save(d,d.revision))')
        self.page.locator('#visual-plan').tap();self.assertIn('catalogue a changé',self.page.locator('#notice').text_content())
        self.assertNotIn('visual_MAIN',self.page.evaluate('CatalogueStore.read().then(d=>d.products[0].values)'))
        self.page.set_viewport_size({'width':412,'height':915});self.assertFalse(self.page.evaluate('document.documentElement.scrollWidth>innerWidth'))
    def test_marketplace_column_aliases(self):
        result=self.page.evaluate("() => ['main_image_url','other_image_url1','MAIN','PT01','image_url_1','image_url_2','Image principale','Image 3'].map(Visuals.guess)")
        self.assertEqual(result,['visual_MAIN','visual_PT01','visual_MAIN','visual_PT01','visual_MAIN','visual_PT01','visual_MAIN','visual_PT02'])
    def test_v1_visual_mapping_export_and_admin_url_rejection(self):
        self.nav('exports')
        catalog=Workbook();sheet=catalog.active;sheet.append(['SKU','Visuel MAIN · URL']);sheet.append(['FICTIF-001','https://example.invalid/public.jpg']);sheet.append(['FICTIF-002','https://semin-produits.onebase.fr/ADMIN/image.jpg'])
        template=Workbook();template.active.append(['SKU','main_image_url'])
        for role,book in [('catalog',catalog),('template',template)]:
            path=Path(self.temp.name)/(role+'-visuals.xlsx');book.save(path);self.page.locator('#'+role+'-file').set_input_files(path);self.page.wait_for_selector('#'+role+'-name:text-is("'+path.name+'")')
        self.page.locator('#start-row').fill('2');self.page.locator('#prepare').tap();self.page.wait_for_selector('#stat-products:text-is("2")')
        self.assertEqual(self.page.get_by_label('Source pour main_image_url',exact=True).input_value(),'2')
        self.page.locator('#select-all').tap();self.page.locator('#check').tap();self.page.wait_for_selector('#stat-errors:text-is("1")');self.assertTrue(self.page.locator('#export').is_disabled())
        self.page.get_by_role('checkbox',name='Sélectionner la ligne 3',exact=True).uncheck();self.page.locator('#check').tap();self.page.wait_for_selector('#stat-errors:text-is("0")');self.page.locator('#export-consent').check()
        with self.page.expect_download() as event:self.page.locator('#export').tap()
        path=Path(self.temp.name)/'visuals-v1-export.xlsx';event.value.save_as(path)
        self.assertEqual(load_workbook(path).active['B2'].value,'https://example.invalid/public.jpg')
