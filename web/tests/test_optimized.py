"""V3.3 public fixtures are synthetic. Real source workbook stays outside Git."""
import unittest
from pathlib import Path
from openpyxl import Workbook,load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
import test_onebase_castorama as base

HEADERS=['Titre/Désignation','EAN','SKU','PVC','STOCK','Image 1/principale/MAIN','Image 2/secondaire/PT01','Image 3/secondaire/PT01','Image 4/secondaire/PT01','Image 5/secondaire/PT01','Image 6/secondaire/PT01','Poids ','Conditionnement ',"Fabriqué/Pays d'origine",*(['PUCE/avantages marketing']*5),'DESCRIPTION']
ROW=[' TITRE FICTIF — texte conservé ','5901234123457','FICTIF-001','PVC INCONNU','10',*(['']*6),'5 Kg','1','Pays fictif',*['Argument '+str(i)+'\nligne conservée' for i in range(1,6)],'<p><strong>Description fictive</strong></p>\n<ul><li>Texte intégral</li></ul>']
class OptimizedTests(unittest.TestCase):
    setUpClass=classmethod(base.OneBaseCastoramaTests.setUpClass.__func__)
    tearDownClass=classmethod(base.OneBaseCastoramaTests.tearDownClass.__func__)
    setUp=base.OneBaseCastoramaTests.setUp
    tearDown=base.OneBaseCastoramaTests.tearDown
    nav=base.OneBaseCastoramaTests.nav
    product=base.OneBaseCastoramaTests.product
    analyse=base.OneBaseCastoramaTests.analyse
    source=base.OneBaseCastoramaTests.source
    save=base.OneBaseCastoramaTests.save
    template=base.OneBaseCastoramaTests.template
    upload=base.OneBaseCastoramaTests.upload
    check=base.OneBaseCastoramaTests.check
    download=base.OneBaseCastoramaTests.download
    def workbook(self,rows=None,name='optimized-fictif.xlsx'):
        w=Workbook();w.active.title='Fiches';w.active.append(HEADERS)
        for row in rows or [ROW]:w.active.append(row)
        path=Path(self.temp.name)/name;w.save(path);return path
    def import_preview(self,path):
        self.nav('optimized');self.page.locator('#opt-file').set_input_files(path);self.page.wait_for_selector('#opt-file-name:text-is("'+path.name+'")');self.page.locator('#opt-analyse').tap();self.page.wait_for_selector('#opt-preview:not([hidden])')
    def confirm(self):
        self.page.locator('#opt-consent').check();self.page.locator('#opt-save').tap();self.page.wait_for_function("() => document.querySelector('#notice').textContent==='Import Excel optimisé confirmé et enregistré localement.'")
    def test_full_text_empty_images_commercial_packaging_and_backup(self):
        self.import_preview(self.workbook());self.assertTrue(self.page.locator('#opt-save').is_disabled());self.assertEqual(self.page.evaluate('CatalogueStore.read().then(d=>d.products.length)'),0)
        self.confirm();data=self.page.evaluate('CatalogueStore.read()');p=data['products'][0]
        self.assertEqual(p['values']['commercialTitle'],ROW[0]);self.assertEqual(p['values']['longDescription'],ROW[19]);self.assertEqual(p['values']['commercialPackaging'],'1');self.assertEqual(p['values']['optimizedWeightSource'],'5 Kg');self.assertEqual(p['values']['optimizedPVC'],'PVC INCONNU')
        for k in ['price','units','weight','lotContents','lotTotalWeight','visual_MAIN']:self.assertNotIn(k,p['values'])
        self.assertEqual(p['optimizedExcel'][0]['source'],ROW)
        self.assertEqual([p['optimizedExcel'][0]['content']['argument'+str(i)] for i in range(1,6)],ROW[14:19])
        self.assertEqual(p['status'],'À compléter');self.page.reload();self.assertEqual(self.page.evaluate('CatalogueStore.read()'),data)
        restored=self.page.evaluate('async()=>{const d=await CatalogueStore.read();return CatalogueModel.restore(CatalogueModel.empty(),CatalogueModel.backup(d))}')
        self.assertEqual(restored['products'],data['products']);self.assertEqual(self.requests,[self.url,self.url])
    def test_updates_require_selection_and_blank_cells_never_erase(self):
        self.product();before=self.page.evaluate('CatalogueStore.read()');row=ROW[:];self.import_preview(self.workbook([row]))
        self.assertFalse(self.page.locator('[data-optimized-row="2"]').is_checked());self.assertEqual(self.page.evaluate('CatalogueStore.read()'),before)
        self.page.locator('[data-optimized-row="2"]').check();self.confirm();first=self.page.evaluate('CatalogueStore.read()');self.assertEqual(first['products'][0]['dossier'],before['products'][0]['dossier'])
        empty=['']*20;empty[2]='FICTIF-001';empty[4]='0';self.import_preview(self.workbook([empty],'optimized-blanks.xlsx'));self.page.locator('[data-optimized-row="2"]').check();self.confirm()
        p=self.page.evaluate('CatalogueStore.read().then(d=>d.products[0])');self.assertEqual(p['values']['commercialTitle'],ROW[0]);self.assertEqual(p['values']['ean'],ROW[1]);self.assertEqual(p['values']['stock'],'0');self.assertEqual(p['values']['longDescription'],ROW[19]);self.assertEqual(len(p['optimizedExcel']),2);self.assertEqual(p['optimizedExcel'][1]['previousValues'],first['products'][0]['values'])
        self.assertEqual(p['dossier'],before['products'][0]['dossier'])
    def test_image_positions_and_multiple_references(self):
        row=ROW[:];row[5:11]=['https://example.invalid/'+str(i)+'.jpg' for i in range(6)]
        other=ROW[:];other[1]='4006381333931';other[2]='FICTIF-002'
        self.import_preview(self.workbook([row,other]));self.confirm();data=self.page.evaluate('CatalogueStore.read()');self.assertEqual(len(data['products']),2)
        for i,slot in enumerate(['MAIN','PT01','PT02','PT03','PT04','PT05']):self.assertEqual(data['products'][0]['values']['visual_'+slot],row[5+i])
        self.assertFalse(any(r.startswith('https://example.invalid') for r in self.requests))
    def test_duplicates_missing_sku_ean_conflicts_and_invalid_urls(self):
        rows=[ROW[:],ROW[:]];self.import_preview(self.workbook(rows));self.assertIn('2 Erreur',self.page.locator('#opt-summary').text_content());self.assertTrue(self.page.locator('[data-optimized-row="2"]').is_disabled())
        self.product();different=ROW[:];different[1]='4006381333931';missing=ROW[:];missing[2]='';bad=ROW[:];bad[2]='FICTIF-003';bad[1]='';bad[5]='https://semin-produits.onebase.fr/ADMIN/image.jpg'
        self.import_preview(self.workbook([different,missing,bad],'optimized-errors.xlsx'));self.assertIn('3 Erreur',self.page.locator('#opt-summary').text_content());self.assertEqual(self.page.evaluate('CatalogueStore.read().then(d=>d.products.length)'),1)
    def test_condition_codes_channels_and_export_html(self):
        self.product();self.import_preview(self.workbook());self.page.locator('[data-optimized-row="2"]').check();self.confirm()
        path=self.template(large=False);w=load_workbook(path);w['Data']['J1']='Etat article';dv=DataValidation(type='list',formula1='"Neuf,Occasion"');dv.add('J2:J200');w['Data'].add_data_validation(dv);w.save(path);self.upload(path)
        self.assertEqual(self.page.get_by_label('Mapping Castorama · Etat article',exact=True).input_value(),'excel.conditionCastorama');self.check();result=load_workbook(self.download());self.assertEqual(result['Data']['J3'].value,'Neuf');self.assertEqual(result['Data']['C3'].value,ROW[0]);self.assertEqual(result['Data']['D3'].value,ROW[19]);self.assertIsNone(result['Data']['G3'].value)
        self.page.locator('#cast-profile-name').fill('Canal fictif Castorama');self.page.locator('#cast-save-profile').tap();self.page.wait_for_function("() => document.querySelector('#notice').textContent.includes('Mapping Castorama enregistré')")
        self.assertEqual(self.page.evaluate("JSON.parse(localStorage.getItem('semin.castorama.local.profiles.v1'))['Canal fictif Castorama'].channel"),'Castorama')
        self.page.locator('#cast-channel').select_option('Amazon');self.page.locator('#cast-analyse').tap();self.check();self.assertIn('liste autorisée',self.page.locator('#cast-errors').text_content())
        self.assertEqual(self.page.evaluate("CatalogueStore.read().then(d=>ProductWorkflow.value(d.products[0],'excel.conditionAmazon'))"),'Nouveau')
        self.assertEqual(self.page.evaluate("CatalogueStore.read().then(d=>ProductWorkflow.value(d.products[0],'excel.conditionADEO'))"),'Neuf')
        self.page.locator('#cast-profile').select_option('Canal fictif Castorama');self.assertIn('autre canal',self.page.locator('#notice').text_content())
        unknown=Workbook();unknown.active.append(['SKU','Etat article']);path=Path(self.temp.name)/'etat-non-verifiable-fictif.xlsx';unknown.save(path);self.upload(path);self.check();self.assertIn('codes techniques non vérifiables',self.page.locator('#cast-errors').text_content())
        coded=DataValidation(type='list',formula1='"CODE-FICTIF-NEUF,CODE-FICTIF-OCCASION"');coded.add('B2:B100');unknown.active.add_data_validation(coded);path=Path(self.temp.name)/'etat-codes-fictifs.xlsx';unknown.save(path);self.upload(path);self.check();self.assertIn('liste autorisée',self.page.locator('#cast-errors').text_content())
    def test_bad_headers_and_formulas_are_rejected_mobile_and_stale_plan(self):
        path=self.workbook();w=load_workbook(path);w.active['D2']='=1+1';w.save(path);self.nav('optimized');self.page.locator('#opt-file').set_input_files(path);self.page.wait_for_selector('#opt-settings:not([hidden])');self.page.locator('#opt-analyse').tap();self.page.wait_for_function("() => document.querySelector('#notice').textContent.includes('Formule non acceptée')")
        self.import_preview(self.workbook(name='optimized-valid.xlsx'));self.page.set_viewport_size({'width':412,'height':915});self.assertFalse(self.page.evaluate('document.documentElement.scrollWidth>innerWidth'))
        self.page.evaluate('CatalogueStore.read().then(d=>CatalogueStore.save(d,d.revision))');self.assertTrue(self.page.locator('#opt-preview').is_hidden());self.assertEqual(self.page.evaluate('CatalogueStore.read().then(d=>d.products.length)'),0)
    def test_v32_database_upgrade_preserves_data_and_blocks_old_writer(self):
        self.product();expected=self.page.evaluate('CatalogueStore.read()')
        context=self.browser.new_context();p=context.new_page();p.on('pageerror',lambda e:self.errors.append(str(e)))
        url=self.url.replace('127.0.0.1','localhost');p.goto(url.rsplit('/',1)[0]+'/GUIDE-VISUELS-URL.md')
        p.evaluate('''data=>new Promise((resolve,reject)=>{const r=indexedDB.open('semin-marketplace-studio',2);r.onupgradeneeded=()=>r.result.createObjectStore('workspace');r.onerror=()=>reject(r.error);r.onsuccess=()=>{const db=r.result,tx=db.transaction('workspace','readwrite');tx.objectStore('workspace').put(data,'catalogue');tx.oncomplete=()=>{db.close();resolve()}}})''',expected)
        p.goto(url);self.assertEqual(p.evaluate('CatalogueStore.read()'),expected)
        result=p.evaluate('''() => new Promise(resolve=>{const r=indexedDB.open('semin-marketplace-studio',2);r.onerror=()=>resolve(r.error.name);r.onsuccess=()=>{r.result.close();resolve('unexpected')}})''')
        self.assertEqual(result,'VersionError');self.assertEqual(p.evaluate('CatalogueStore.read()'),expected);context.close()
    def test_corrupt_archive_is_rejected_without_altering_catalogue_desktop(self):
        self.import_preview(self.workbook());self.confirm();before=self.page.evaluate('CatalogueStore.read()')
        blocked=self.page.evaluate('''async()=>{const d=await CatalogueStore.read();d.products[0].optimizedExcel[0].version=999;try{await CatalogueStore.save(d,d.revision);return false}catch{return true}}''')
        self.assertTrue(blocked);self.assertEqual(self.page.evaluate('CatalogueStore.read()'),before)
        self.page.set_viewport_size({'width':1280,'height':800});self.assertFalse(self.page.evaluate('document.documentElement.scrollWidth>innerWidth'))
