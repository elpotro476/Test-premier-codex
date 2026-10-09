"""Public tests use synthetic text and workbooks ONLY. Never add a real catalogue here."""
import json
from pathlib import Path
from copy import copy
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from openpyxl import Workbook,load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill
import unittest
import test_browser as base

TEXT="""PEINTURE FICTIVE MAT
*Code produit*
000000000-00001
*Logo marque*
[image](https://example.invalid/fiction.png?key=INVALID)
*Marque*
Marque FICTIVE - FICTIVE
*Catégorie*
Peinture
Argumentaire
*Définition technique*
Peinture acrylique de démonstration finition mate
intérieure et extérieure
*Les + produits 1*
Argument fictif un
*Les + produits 2*
Argument fictif deux
*Les + produits 3*
Argument fictif trois
*Les + produits 4*
Argument fictif quatre
Infos complémentaires
*Support admis*
Murs et plafonds.
*Conditionnement*
Seau fictif
*Temps de séchage*
Sec : 2 h.
Recouvrable : 6 h.
*Résistance à l'abrasion humide*
Norme fictive : classe 2.
Référence article : FICTIF-001"""

class OneBaseCastoramaTests(unittest.TestCase):
    setUpClass=classmethod(base.BrowserTests.setUpClass.__func__)
    tearDownClass=classmethod(base.BrowserTests.tearDownClass.__func__)
    setUp=base.BrowserTests.setUp
    tearDown=base.BrowserTests.tearDown
    def nav(self,name):self.page.locator(f'[data-studio-view="{name}"]').tap()
    def analyse(self,text=TEXT):
        self.nav('onebase');self.page.locator('#ob-raw').fill(text);self.page.locator('#ob-analyse').tap();self.page.wait_for_selector('#ob-source-fields [data-source="sku"]')
    def source(self,field):return self.page.locator(f'#ob-source-fields [data-source="{field}"]')
    def save(self):
        self.page.locator('#ob-preview').tap();self.page.wait_for_selector('#ob-confirm-panel:not([hidden])');self.page.locator('#ob-consent').check();self.page.locator('#ob-confirm').tap();self.page.wait_for_selector('#ob-confirm-panel',state='hidden');self.page.wait_for_function("() => document.querySelector('#notice').textContent.includes('enregistrée localement')")
    def product(self,ean=True):
        self.analyse()
        if ean:self.source('ean').fill('5901234123457')
        self.page.locator('#ob-suggest').tap();self.save()
    def template(self,validation='list',large=True,metadata=False):
        wb=Workbook();s=wb.active;s.title='Data'
        labels=['Boutique SKU','EAN','Nom','Texte description','Puces avantages marketing 1','Marque','Prix Offre','Quantité Offre','Contenance (ml)']
        s.append(labels);s.append(['FICTIF-OLD','4006381333931','Ancien produit fictif','<p>Fictif</p>','Fictif','FICTIVE',1.25,9,1000])
        for c in s[2]:c.fill=PatternFill('solid',fgColor='EFF5EE')
        s['G2'].number_format='0.00';s['A1'].comment=Comment('SKU obligatoire','Fictif');s.freeze_panes='A2';s.column_dimensions['C'].width=36
        refs=wb.create_sheet('ReferenceData');refs.append(['Marques fictives']);refs.append(['FICTIVE']);refs.cell(28000 if large else 3,1,'AUTRE FICTIVE')
        wb.defined_names.add(DefinedName('ListeMarquesFictives',attr_text="ReferenceData!$A$2:$A$"+str(28000 if large else 3)))
        dv=DataValidation(type=validation,formula1='ListeMarquesFictives' if validation=='list' else 'AND(F3<>"",LEN(F3)>2)',allow_blank=True);dv.add('F2:F200');s.add_data_validation(dv)
        price=DataValidation(type='decimal',operator='between',formula1='0',formula2='1000',allow_blank=True);price.add('G2:G200');s.add_data_validation(price)
        instructions=wb.create_sheet('Instructions');instructions['A1']='Exemple fictif uniquement';instructions['B2']='=1+1'
        if metadata:
            cols=wb.create_sheet('Columns');cols.append(['Code','Libellé','Description','Exemple','Catégorie fictive']);cols.append(['ean','EAN','Fictif',None,'REQUIRED']);cols.append(['price','Prix Offre','Fictif',None,'REQUIRED'])
        path=Path(self.temp.name)/'template-fictif.xlsx';wb.save(path);return path
    def upload(self,path):
        self.nav('castorama');self.page.locator('#cast-file').set_input_files(path);self.page.wait_for_selector('#cast-file-name:text-is("'+path.name+'")');self.page.locator('#cast-analyse').tap();self.page.wait_for_selector('#cast-prepared:not([hidden])')
    def check(self):
        self.page.locator('#cast-check').tap();self.page.wait_for_function("() => document.querySelector('#cast-review-message').textContent.includes('Contrôle réussi') || document.querySelector('#cast-review-message').textContent.includes('erreur(s)')")
    def download(self):
        with self.page.expect_download() as e:self.page.locator('#cast-export').tap()
        path=Path(self.temp.name)/'export-fictif.xlsx';e.value.save_as(path);return path
    def test_extraction_multiline_missing_values_four_arguments_no_remote_calls(self):
        self.analyse();self.assertEqual(self.source('sku').input_value(),'FICTIF-001');self.assertEqual(self.source('onebaseRef').input_value(),'000000000-00001');self.assertEqual(self.source('brand').input_value(),'FICTIVE');self.assertEqual(self.source('designation').input_value(),'PEINTURE FICTIVE MAT');self.assertEqual(self.source('drying').input_value(),'2 h');self.assertEqual(self.source('recoat').input_value(),'6 h.')
        for field in ['ean','capacity','coverage','price','stock']:self.assertEqual(self.source(field).input_value(),'')
        self.page.locator('#ob-suggest').tap();self.assertEqual(self.page.get_by_label('Enrichissement · Argument commercial 4',exact=True).input_value(),'Argument fictif quatre');self.assertEqual(self.page.get_by_label('Enrichissement · Argument commercial 5',exact=True).input_value(),'');self.assertEqual(self.requests,[self.url])
        parsed=self.page.evaluate("OneBase.parse('Référence OneBase : 000000000-00001\\nRéférence article : FICTIF-A\\nRéférence article : FICTIF-B')");self.assertEqual(parsed['extracted']['sku'],'');self.assertTrue(parsed['warnings'])
    def test_structured_record_original_source_overrides_and_persistence(self):
        self.analyse();self.page.locator('#ob-suggest').tap();self.source('designation').fill('Désignation fictive corrigée');self.page.get_by_label('Personnaliser Castorama · title',exact=True).check();self.page.get_by_label('Castorama · Titre optimisé',exact=True).fill('Titre Castorama fictif');self.save()
        data=self.page.evaluate('CatalogueStore.read()');p=data['products'][0];self.assertEqual(p['dossier']['source']['raw'],TEXT);self.assertEqual(p['dossier']['source']['extracted']['designation'],'PEINTURE FICTIVE MAT');self.assertEqual(p['values']['designation'],'Désignation fictive corrigée');self.assertEqual(p['dossier']['castorama']['title'],'Titre Castorama fictif');self.assertNotEqual(p['dossier']['enrichment']['title'],p['dossier']['castorama']['title'])
        self.page.reload();self.nav('onebase');self.page.locator('#ob-existing').select_option(p['id']);self.assertEqual(self.page.locator('#ob-raw').input_value(),TEXT);self.assertEqual(self.source('designation').input_value(),'Désignation fictive corrigée');self.assertEqual(self.page.get_by_label('Castorama · Titre optimisé',exact=True).input_value(),'Titre Castorama fictif')
    def test_existing_sku_requires_confirmation_and_stale_update_is_refused(self):
        self.product();self.page.locator('#ob-new').tap();self.page.locator('#ob-raw').fill(TEXT.replace('PEINTURE FICTIVE MAT','NOUVELLE DÉSIGNATION FICTIVE'));self.page.locator('#ob-analyse').tap();self.page.locator('#ob-preview').tap();self.page.wait_for_selector('#ob-confirm-panel:not([hidden])');self.assertTrue(self.page.locator('#ob-confirm').is_disabled());self.assertIn('5901234123457',self.page.locator('#ob-changes').text_content());self.assertEqual(self.page.evaluate('CatalogueStore.read().then(d=>d.products.length)'),1)
        self.page.evaluate('CatalogueStore.read().then(d=>CatalogueStore.save(d,d.revision))');self.page.locator('#ob-consent').check();self.page.locator('#ob-confirm').tap();self.page.wait_for_function("() => document.querySelector('#notice').className==='error'");self.assertEqual(self.page.evaluate('CatalogueStore.read().then(d=>d.products[0].values.designation)'),'PEINTURE FICTIVE MAT')
    def test_old_v2_backup_is_supported_and_bad_dossier_rejected(self):
        result=self.page.evaluate("(()=>{const d=CatalogueModel.empty(); const backup=CatalogueModel.backup(d);const restored=CatalogueModel.restore(d,backup);return restored.schemaVersion})()");self.assertEqual(result,1)
        self.product();self.assertTrue(self.page.evaluate("CatalogueStore.read().then(d=>{d.products[0].dossier.source.raw=42;try{CatalogueModel.validateSnapshot(d);return false}catch{return true}})"))
        backup=self.page.evaluate('CatalogueStore.read().then(d=>CatalogueModel.backup(d))');self.assertEqual(backup['catalogue']['products'][0]['dossier']['source']['raw'],TEXT)
    def test_filled_template_large_reference_export_preserves_parts_rows_and_rules(self):
        self.product();template=self.template();self.upload(template);self.assertEqual(self.page.locator('#cast-start').input_value(),'3');self.assertEqual(self.page.get_by_label('Mapping Castorama · Boutique SKU',exact=True).input_value(),'master.sku');self.check();self.assertFalse(self.page.locator('#cast-export').is_disabled());path=self.download();old=load_workbook(template);new=load_workbook(path)
        self.assertEqual(old.sheetnames,new.sheetnames);self.assertEqual(list(old['Data'].values)[1],list(new['Data'].values)[1]);self.assertEqual(new['Data']['A3'].value,'FICTIF-001');self.assertEqual(new['Data']['B3'].value,'5901234123457');self.assertEqual(new['Data']['G3'].value,None);self.assertEqual(new['Data']['H3'].value,None);self.assertEqual(new['Data']['I3'].value,None);self.assertEqual(new['Data'].freeze_panes,old['Data'].freeze_panes);self.assertEqual(copy(new['Data']['A3'].fill),copy(old['Data']['A2'].fill));self.assertEqual(new['Data']['A1'].comment.text,old['Data']['A1'].comment.text);self.assertEqual(new['Data'].data_validations.to_tree().attrib,old['Data'].data_validations.to_tree().attrib)
        self.assertEqual(ET.tostring(new['Data'].data_validations.to_tree()),ET.tostring(old['Data'].data_validations.to_tree()))
        with ZipFile(template) as a,ZipFile(path) as b:self.assertEqual(set(a.namelist()),set(b.namelist()));self.assertEqual([n for n in a.namelist() if a.read(n)!=b.read(n)],['xl/worksheets/sheet1.xml'])
        self.assertEqual(self.requests,[self.url])
    def test_real_requirements_metadata_and_missing_values_block_export(self):
        self.product(ean=False);self.upload(self.template(metadata=True));self.check();self.assertTrue(self.page.locator('#cast-export').is_disabled());errors=self.page.locator('#cast-errors').text_content();self.assertIn('EAN',errors);self.assertIn('Prix Offre',errors);self.assertIn('obligatoire',errors)
    def test_manual_values_are_per_product_and_mapping_backup_has_no_values(self):
        self.product();self.upload(self.template());self.page.get_by_label('Mapping Castorama · Prix Offre',exact=True).select_option('manual');self.page.get_by_label('Valeur Castorama · Prix Offre',exact=True).fill('8,50');self.check();self.assertIn('Enregistrez les saisies',self.page.locator('#cast-errors').text_content());self.page.locator('#cast-values-consent').check();self.page.locator('#cast-save-values').tap();self.page.wait_for_function("() => document.querySelector('#notice').textContent.includes('Valeurs Castorama enregistrées')");self.check();self.assertFalse(self.page.locator('#cast-export').is_disabled());self.assertEqual(load_workbook(self.download())['Data']['G3'].value,8.5)
        self.page.locator('#cast-save-profile').tap();self.page.wait_for_function("() => document.querySelector('#notice').textContent.includes('Mapping Castorama enregistré')")
        with self.page.expect_download() as e:self.page.locator('#cast-backup-profiles').tap()
        path=Path(self.temp.name)/'cast-mappings.json';e.value.save_as(path);content=path.read_text();self.assertNotIn('FICTIF-001',content);self.assertNotIn('8,50',content);self.assertNotIn('5901234123457',content)
        self.page.reload();self.upload(self.template());self.page.locator('#cast-profile').select_option('Castorama');self.assertEqual(self.page.get_by_label('Mapping Castorama · Prix Offre',exact=True).input_value(),'manual')
    def test_occupied_merged_duplicate_and_invalid_list_destination_rejected(self):
        self.product();self.upload(self.template());self.page.locator('#cast-start').fill('2');self.check();self.assertIn('déjà une valeur',self.page.locator('#cast-errors').text_content());self.page.locator('#cast-start').fill('3');self.page.get_by_label('Mapping Castorama · Marque',exact=True).select_option('manual');self.page.get_by_label('Valeur Castorama · Marque',exact=True).fill('MARQUE NON AUTORISÉE');self.page.locator('#cast-values-consent').check();self.page.locator('#cast-save-values').tap();self.page.wait_for_function("() => document.querySelector('#notice').textContent.includes('Valeurs Castorama enregistrées')");self.check();self.assertIn('liste autorisée',self.page.locator('#cast-errors').text_content())
        p=self.template();w=load_workbook(p);w['Data'].merge_cells('A3:B3');w.save(p);self.upload(p);self.check();self.assertIn('fusionnée',self.page.locator('#cast-errors').text_content())
    def test_unknown_validations_outside_range_and_html_injection_blocked(self):
        self.product();self.upload(self.template(validation='custom'));self.check();self.assertIn('custom',self.page.locator('#cast-errors').text_content());self.upload(self.template());self.page.locator('#cast-start').fill('201');self.check();self.assertIn('hors de la plage',self.page.locator('#cast-errors').text_content())
        self.nav('onebase');p=self.page.evaluate('CatalogueStore.read().then(d=>d.products[0].id)');self.page.locator('#ob-existing').select_option(p);self.page.get_by_label('Enrichissement · Description HTML',exact=True).fill('<img src="https://example.invalid/attack.png" onerror="alert(1)">');self.save();self.nav('castorama');self.page.locator('#cast-start').fill('3');self.check();self.assertIn('HTML',self.page.locator('#cast-errors').text_content());self.assertEqual(self.requests,[self.url])
    def test_multiple_volumes_and_category_inconsistency_require_review(self):
        text=TEXT.replace('Seau fictif','Seaux de 4 et 15 L').replace('*Catégorie*\nPeinture','*Catégorie*\nEnduit');self.analyse(text);self.assertIn('Plusieurs contenances',self.page.locator('#ob-warnings').text_content());self.assertIn('Enduit',self.page.locator('#ob-warnings').text_content());self.assertEqual(self.source('capacity').input_value(),'');self.save();self.upload(self.template());self.check();self.assertIn('Plusieurs contenances',self.page.locator('#cast-errors').text_content());self.assertIn('Enduit',self.page.locator('#cast-errors').text_content())
    def test_touch_layout_no_overflow_and_local_mapping_storage_failure(self):
        self.product();self.upload(self.template());self.page.set_viewport_size({'width':412,'height':915});self.assertFalse(self.page.evaluate('document.documentElement.scrollWidth>innerWidth'));self.page.evaluate("localStorage.setItem('semin.castorama.local.profiles.v1','invalid')");self.page.locator('#cast-save-profile').tap();self.page.wait_for_function("() => document.querySelector('#notice').className==='error'");self.assertEqual(self.page.evaluate("localStorage.getItem('semin.castorama.local.profiles.v1')"),'invalid')
    def test_downloadable_synthetic_demo_runs_without_external_data(self):
        self.nav('onebase');self.page.locator('#ob-demo').tap();self.page.wait_for_function("() => document.querySelector('#ob-source-fields [data-source=sku]').value==='FICTIF-OB-001'");self.page.locator('#ob-suggest').tap();self.save();self.nav('castorama')
        with self.page.expect_download() as event:self.page.locator('#cast-demo-template').tap()
        path=Path(self.temp.name)/'demo-fictif.xlsx';event.value.save_as(path);self.assertIn('FICTIF',load_workbook(path)['Instructions']['A1'].value)
        self.upload(path);self.check();self.assertFalse(self.page.locator('#cast-export').is_disabled());w=load_workbook(self.download());self.assertEqual(w['Data']['A2'].value,'FICTIF-EXISTANT');self.assertEqual(w['Data']['A3'].value,'FICTIF-OB-001');self.assertEqual(self.requests,[self.url])

if __name__=='__main__':unittest.main()
