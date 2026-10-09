"""Excel processing: values are copied only from the supplied catalogue."""
from copy import copy
from io import BytesIO
import re
import unicodedata
from openpyxl import load_workbook

ALIASES = {
    'sku': ['sku', 'reference', 'reference produit', 'seller sku', 'item sku', 'ref'],
    'ean': ['ean', 'ean13', 'gtin', 'barcode', 'code barre', 'product id'],
    'name': ['nom', 'titre', 'designation', 'product name', 'item name', 'title'],
    'description': ['description', 'product description'],
    'price': ['prix', 'prix ttc', 'price', 'standard price'],
    'brand': ['marque', 'brand', 'brand name'],
    'weight': ['poids', 'weight', 'item weight'],
    'stock': ['stock', 'quantite', 'quantity'],
}

def normalize(value):
    text = unicodedata.normalize('NFKD', str(value or '').lower())
    return re.sub(r'[^a-z0-9]+', ' ', ''.join(c for c in text if not unicodedata.combining(c))).strip()

def kind(label):
    label = normalize(label.replace('*', ''))
    return next((key for key, values in ALIASES.items() if label in values), None)

def workbook(data):
    if len(data) > 15 * 1024 * 1024:
        raise ValueError('Fichier trop volumineux (maximum 15 Mo).')
    # Prevent zip bombs before openpyxl decompresses the workbook.
    from zipfile import ZipFile
    with ZipFile(BytesIO(data)) as archive:
        if sum(i.file_size for i in archive.infolist()) > 80 * 1024 * 1024:
            raise ValueError('Classeur trop volumineux après décompression.')
    return load_workbook(BytesIO(data), data_only=False)

def inspect(data):
    wb = workbook(data)
    result = []
    for ws in wb:
        if ws.max_row > 50000 or ws.max_column > 250:
            raise ValueError('Limite de démonstration : 50 000 lignes et 250 colonnes par feuille.')
        preview = [[str(v) if v is not None else '' for v in row] for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 12), values_only=True)]
        result.append({'name': ws.title, 'rows': ws.max_row, 'preview': preview})
    return result

def columns(ws, header):
    if not 1 <= header <= ws.max_row:
        raise ValueError('Ligne d’en-tête invalide.')
    result = []
    for cell in ws[header]:
        if cell.value is not None:
            label = str(cell.value)
            comment = cell.comment.text.lower() if cell.comment else ''
            result.append({'id': str(cell.column), 'label': label, 'kind': kind(label),
                           'required': '*' in label or 'obligatoire' in comment or 'required' in comment})
    if not result:
        raise ValueError('La ligne choisie ne contient pas d’en-têtes.')
    return result

def read_catalog(data, sheet, header):
    ws = workbook(data)[sheet]
    cols = columns(ws, header)
    rows = []
    for r in range(header + 1, ws.max_row + 1):
        values = {}
        for col in cols:
            cell = ws.cell(r, int(col['id']))
            value = cell.value
            if cell.data_type == 'f':
                raise ValueError(f'Catalogue : formule en {cell.coordinate}. Fournissez des valeurs figées.')
            if col['kind'] in ('sku', 'ean') and isinstance(value, (int, float)):
                if isinstance(value, float) and not value.is_integer():
                    raise ValueError(f'Identifiant non entier en {cell.coordinate}.')
                value = str(int(value))
                if re.fullmatch(r'0+', cell.number_format):
                    value = value.zfill(len(cell.number_format))
            values[col['id']] = value
        if any(v is not None and v != '' for v in values.values()):
            rows.append({'id': r, 'values': values})
    if not rows:
        raise ValueError('Le catalogue ne contient aucun produit.')
    return cols, rows

def suggest(source, target):
    mapping = {}
    for dest in target:
        matches = [c for c in source if normalize(c['label']) == normalize(dest['label'].replace('*', ''))]
        if not matches and dest['kind']:
            matches = [c for c in source if c['kind'] == dest['kind']]
        mapping[dest['id']] = matches[0]['id'] if len(matches) == 1 else ''
    return mapping

def ean_valid(value):
    s = str(value)
    if not re.fullmatch(r'\d{13}', s):
        return False
    return (sum(int(c) * (1 if i % 2 == 0 else 3) for i, c in enumerate(s[:12])) + int(s[-1])) % 10 == 0

def check(rows, target, mapping, required, selected):
    chosen = [r for r in rows if r['id'] in selected]
    if not chosen:
        raise ValueError('Sélectionnez au moins un produit.')
    errors, output = [], []
    seen = {}
    for row in chosen:
        values = {}
        for col in target:
            source = mapping.get(col['id'])
            value = row['values'].get(source) if source else None
            values[col['id']] = value
            message = None
            if value is None or value == '':
                if col['id'] in required:
                    message = 'Champ obligatoire manquant'
            elif isinstance(value, str) and value.startswith(('=', '+', '-', '@')):
                message = 'Valeur ambiguë pouvant être interprétée comme une formule'
            elif col['kind'] == 'ean' and not ean_valid(value):
                message = 'EAN-13 invalide (13 chiffres et clé de contrôle)'
            elif col['kind'] in ('price', 'weight', 'stock'):
                try:
                    number = float(str(value).replace(',', '.'))
                    import math
                    if not math.isfinite(number) or number < 0 or (col['kind'] == 'stock' and not number.is_integer()):
                        raise ValueError()
                    values[col['id']] = int(number) if col['kind'] == 'stock' else number
                except (ValueError, TypeError):
                    message = 'Nombre positif ou nul requis' + (' et entier' if col['kind'] == 'stock' else '')
            if value is not None and value != '' and col['kind'] in ('sku', 'ean'):
                key = (col['id'], str(value))
                if key in seen:
                    message = f'Identifiant dupliqué (ligne {seen[key]})'
                else:
                    seen[key] = row['id']
            if message:
                errors.append({'row': row['id'], 'column': col['label'], 'message': message})
        output.append({'id': row['id'], 'values': values})
    return {'rows': output, 'errors': errors, 'count': len(output)}

def export(data, sheet, header, start, target, result):
    if result['errors']:
        raise ValueError('Corrigez les erreurs avant export.')
    if start <= header:
        raise ValueError('La première ligne de données doit suivre les en-têtes.')
    wb = workbook(data)
    ws = wb[sheet]
    last = start + len(result['rows']) - 1
    if last > 50000:
        raise ValueError('Export limité à 50 000 lignes.')
    mapped_cols = {int(c['id']) for c in target}
    for merged in ws.merged_cells.ranges:
        if merged.min_row <= last and merged.max_row >= start and any(merged.min_col <= c <= merged.max_col for c in mapped_cols):
            raise ValueError('Les cellules de destination sont fusionnées. Choisissez une zone de données non fusionnée.')
    # Never erase instructions, defaults or formulas present in destination cells.
    for r in range(start, last + 1):
        for c in mapped_cols:
            if ws.cell(r, c).value is not None:
                raise ValueError(f'La cellule {ws.cell(r, c).coordinate} contient déjà une valeur. Choisissez une zone vierge.')
    for offset, row in enumerate(result['rows']):
        r = start + offset
        if r != start:
            ws.row_dimensions[r].height = ws.row_dimensions[start].height
        for col in target:
            c = int(col['id'])
            dest = ws.cell(r, c)
            model = ws.cell(start, c)
            if r != start and not dest.has_style:
                dest._style = copy(model._style)
            value = row['values'][col['id']]
            dest.value = value
            if col['kind'] in ('sku', 'ean') and value is not None:
                dest.value = str(value)
            if isinstance(value, str):
                dest.data_type = 's'
    # Extend validation ranges only when the model data row has that validation.
    from openpyxl.utils import get_column_letter
    for validation in ws.data_validations.dataValidation:
        for c in mapped_cols:
            letter = get_column_letter(c)
            if f'{letter}{start}' in validation:
                validation.add(f'{letter}{start}:{letter}{last}')
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
